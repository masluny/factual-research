"""Tables: the agent fills evidence tables cell by cell (value + verbatim
quote + location); every sourced cell is verified against the source
before it is stored. Built-in tables: sources, extraction (evidence table),
summary of findings, claim × source matrix, forest plot, and free-form
custom tables from a JSON spec. Output: md, csv, html, latex, xlsx, json."""
from __future__ import annotations

import csv
import html as H
import io
import json
import math
import re
from dataclasses import dataclass, field
from pathlib import Path

from .evidence import flags, summarize_claim, weight
from .model import DESIGNS
from .project import Plan
from .store import INCLUDED, Store
from .textutil import extract_numbers, find_quote, is_significant, number_in, number_index


# Built-in table text in the language of the final text (plan.toml `lang`).
# A pack field can also carry its own `label_<lang>`.
L10N = {"pl": {
    "Characteristics and results of included studies": "Charakterystyka i wyniki włączonych badań",
    "Study": "Badanie",
    "Every non-empty cell is backed by a verbatim quote from the source (hover in HTML report).":
        "Każda niepusta komórka jest poparta dosłownym cytatem ze źródła (widocznym po najechaniu kursorem w raporcie HTML).",
    "{n} cell(s) failed verification - marked ⚠.": "Komórki, które nie przeszły weryfikacji: {n} (oznaczone ⚠).",
    "Included sources": "Uwzględnione źródła",
    "Year": "Rok", "Venue": "Czasopismo", "Weight": "Waga", "Text": "Tekst", "Flags": "Uwagi",
    "Weight = evidence weight of the design in the active domain pack(s) (0–1).":
        "Waga = waga dowodowa typu badania w aktywnych pakietach dziedzinowych (0–1).",
    "Summary of findings": "Podsumowanie wyników",
    "Claim": "Twierdzenie", "Studies": "Badania", "Designs": "Typy badań", "For / against / mixed": "Za / przeciw / mieszane",
    "Consensus": "Zgodność", "Certainty": "Pewność", "Key effects": "Kluczowe efekty", "Why": "Uzasadnienie",
    "Consensus = design-weighted share of supporting evidence. Certainty = GRADE-style: start from the best design, "
    "capped by the reviews' own GRADE / risk-of-bias rating, downgrade for inconsistency, imprecision, preprint-only, "
    "industry funding.":
        "Zgodność = udział dowodów wspierających, ważony typem badania. Pewność = w stylu GRADE: punktem wyjścia jest "
        "najlepszy typ badania, ograniczony własną oceną GRADE / ryzyka błędu systematycznego w przeglądach, obniżana "
        "za niespójność, nieprecyzyjność, wyłącznie preprinty i finansowanie przez przemysł.",
    "Evidence matrix (claim × source)": "Macierz dowodów (twierdzenie × źródło)",
    "Source": "Źródło", "Claims: ": "Twierdzenia: ",
    "Forest plot": "Wykres typu forest plot", "Effect estimates (95% CI)": "Oszacowania efektu (95% CI)",
    "favours intervention": "na korzyść interwencji", "favours control": "na korzyść kontroli",
    # field labels of the built-in packs
    "Design": "Typ badania", "Design / type": "Typ / rodzaj", "Population": "Populacja", "N": "N",
    "Sample (N, who)": "Próba (N, kto)", "Context / sample": "Kontekst / próba", "Specimens / n": "Próbki / n",
    "Intervention / exposure": "Interwencja / ekspozycja", "Comparator": "Grupa porównawcza",
    "Duration / follow-up": "Czas trwania / obserwacji", "Primary outcome": "Główny punkt końcowy",
    "Effect (95% CI)": "Efekt (95% CI)", "Effect size": "Wielkość efektu", "Adverse events": "Zdarzenia niepożądane",
    "Funding / COI": "Finansowanie / konflikt interesów", "Measures": "Narzędzia pomiaru",
    "Replication status": "Status replikacji", "Limitations": "Ograniczenia",
    "Limitations / external validity": "Ograniczenia / trafność zewnętrzna",
    "Systematics / limitations": "Błędy systematyczne / ograniczenia", "Result": "Wynik",
    "Result (± uncertainty)": "Wynik (± niepewność)", "Key result": "Kluczowy wynik", "Main estimate": "Główne oszacowanie",
    "Method": "Metoda", "Method / model": "Metoda / model", "Identification / method": "Identyfikacja / metoda",
    "Metric": "Miara", "Task": "Zadanie", "Mechanism": "Mechanizm", "Status": "Status", "Data": "Dane",
    "Data / model": "Dane / model", "Data / facility": "Dane / aparatura", "Dataset / benchmark": "Zbiór danych / benchmark",
    "Best baseline": "Najlepszy punkt odniesienia", "Baseline / comparison": "Punkt odniesienia / porównanie",
    "Compute / scale": "Moc obliczeniowa / skala", "Code / data available": "Dostępność kodu / danych",
    "Experiment / model": "Eksperyment / model", "Observable / quantity": "Obserwabla / wielkość",
    "Significance (σ / CL)": "Istotność (σ / CL)", "System / material": "Układ / materiał",
    "Substance / system": "Substancja / układ", "Conditions": "Warunki",
    "Conditions (T, p, solvent)": "Warunki (T, p, rozpuszczalnik)", "Region / system": "Region / system",
    "Period / scenario": "Okres / scenariusz", "Stated confidence / likelihood": "Deklarowana pewność / prawdopodobieństwo",
    "Country / market / period": "Kraj / rynek / okres", "Act / court": "Akt / sąd",
    "Provision / signature": "Przepis / sygnatura", "Rule / holding": "Reguła / teza",
    "Date / version in force": "Data / wersja obowiązująca",
}}


def tr(text: str, lang: str = "en") -> str:
    return L10N.get(lang, {}).get(text, text)


def field_label(f: dict, lang: str = "en") -> str:
    return f.get(f"label_{lang}") or tr(f.get("label", f["id"]), lang)


@dataclass
class Table:
    title: str
    columns: list[str]
    rows: list[list] = field(default_factory=list)       # cells: str or dict(value, key, quote, loc, ok)
    notes: list[str] = field(default_factory=list)
    kind: str = "custom"

    def cell_text(self, c) -> str:
        if isinstance(c, dict):
            return str(c.get("value", ""))
        return "" if c is None else str(c)


# ---------------------------------------------------------------------------
# cell verification
# ---------------------------------------------------------------------------

def check_cell(st: Store, key: str, value: str, quote: str, loc: str = "") -> tuple[float, list[str]]:
    """Return (score 0..1, problems). Quote must be in the source; every
    significant number in the value must be in the quote or the source."""
    problems = []
    depth, full, pages = st.source_text(key)
    score = 1.0
    if quote and len(quote.split()) < 6:
        problems.append("quote is too short to prove anything (<6 words) - use the whole sentence")
        score = min(score, 0.8)
    if quote:
        sim, at = find_quote(quote, full)
        if sim < 0.9:
            problems.append(f"quote not found in @{key} ({depth}, best {sim:.0%})")
            score = min(score, sim * 0.5)
        elif loc:
            from .citations import quote_location
            w = st.get(key)
            ok, where = quote_location(quote, loc, pages, (w.abstract if w else ""))
            if ok is False:
                problems.append(f"quote found at {where or '?'}, not at {loc}")
                score = min(score, 0.9)
    elif value and any(is_significant(n) for n in extract_numbers(value)):
        problems.append("numeric cell without a supporting quote")
        score = min(score, 0.6)
    qidx = number_index(quote) if quote else []
    sidx = number_index(full)
    for n in extract_numbers(value or ""):
        if not is_significant(n):
            continue
        if quote and number_in(n, {v for v, _ in qidx}, qidx) != "missing":
            continue
        if number_in(n, {v for v, _ in sidx}, sidx) != "missing":
            if quote:
                problems.append(f"{n.raw} is in the source but not in the quote")
                score = min(score, 0.85)
            continue
        problems.append(f"{n.raw} not found in @{key}" + ("" if depth == "fulltext" else f" ({depth} only)"))
        score = min(score, 0.3 if depth == "fulltext" else 0.5)
    return round(score, 3), problems


# ---------------------------------------------------------------------------
# extraction (evidence table)
# ---------------------------------------------------------------------------

def extraction_template(st: Store, plan: Plan, keys: list[str] | None = None) -> dict:
    fields = plan.combined().fields()
    keys = keys or [k for k, s, w in st.works(status=INCLUDED)]
    out = {"_instructions": "Fill value + verbatim quote (+ loc like 'p. 4' / 'table 2' / 'sec. Results') for each "
                            "field from the source text (`factual-research show KEY --text` / `factual-research find`). Leave "
                            "value empty if not reported; write 'NR' for not reported. Then: factual-research extract import FILE",
           "_fields": {f["id"]: f.get("label", f["id"]) + (f" - {f['hint']}" if f.get("hint") else "") for f in fields},
           "studies": {}}
    for k in keys:
        w = st.get(k)
        existing = {c["field"]: dict(c) for c in st.cells(k)}
        depth = st.source_text(k)[0]
        out["studies"][k] = {"_ref": f"{(w.authors or ['?'])[0]} {w.year} - {w.title[:100]} [{w.design}; {depth}]",
                             **{f["id"]: {"value": existing.get(f["id"], {}).get("value", ""),
                                          "quote": existing.get(f["id"], {}).get("quote", ""),
                                          "loc": existing.get(f["id"], {}).get("loc", "")}
                                for f in fields}}
    return out


def extract_set(st: Store, key: str, fld: str, value: str, quote: str = "", loc: str = "") -> tuple[float, list[str]]:
    score, problems = (1.0, []) if value.strip().upper() in ("NR", "N/A", "") and not quote else \
        check_cell(st, key, value, quote, loc)
    if quote and not loc:
        from .citations import quote_location
        loc = quote_location(quote, "", st.source_text(key)[2])[1]
    st.set_cell(key, fld, value, quote, loc, score)
    return score, problems


def extract_import(st: Store, data: dict) -> list[tuple[str, str, float, list[str]]]:
    res = []
    studies = data.get("studies", data)
    for key, cells in studies.items():
        if key.startswith("_"):
            continue
        k = st.resolve(key)
        if not k:
            res.append((key, "*", 0.0, ["unknown key"]))
            continue
        for fld, c in cells.items():
            if fld.startswith("_"):
                continue
            if isinstance(c, str):
                c = {"value": c}
            if not (c.get("value") or c.get("quote")):
                continue
            sc, pr = extract_set(st, k, fld, c.get("value", ""), c.get("quote", ""), c.get("loc", ""))
            res.append((k, fld, sc, pr))
    return res


def extraction_table(st: Store, plan: Plan, keys: list[str] | None = None, fields: list[str] | None = None,
                     lang: str = "en") -> Table:
    fdefs = plan.combined().fields()
    if fields:
        fdefs = [f for f in fdefs if f["id"] in fields] + [{"id": f, "label": f} for f in fields
                                                            if f not in {x["id"] for x in fdefs}]
    cells = {}
    for c in st.cells():
        cells.setdefault(c["key"], {})[c["field"]] = c
    keys = keys or [k for k, s, w in st.works(status=INCLUDED) if k in cells]
    t = Table(tr("Characteristics and results of included studies", lang),
              [tr("Study", lang)] + [field_label(f, lang) for f in fdefs], kind="extraction")
    bad = 0
    for k in keys:
        row = [{"value": f"[@{k}]", "key": k, "cite": True}]
        for f in fdefs:
            c = cells.get(k, {}).get(f["id"])
            if c:
                ok = (c["verified"] or 0) >= 0.85
                bad += 0 if ok else 1
                row.append({"value": c["value"], "key": k, "quote": c["quote"], "loc": c["loc"], "ok": ok,
                            "score": c["verified"]})
            else:
                row.append("")
        t.rows.append(row)
    keep = [0] + [j for j in range(1, len(t.columns)) if any(isinstance(r[j], dict) and r[j].get("value") for r in t.rows)]
    if len(keep) < len(t.columns):
        t.columns = [t.columns[j] for j in keep]
        t.rows = [[r[j] for j in keep] for r in t.rows]
    t.notes.append(tr("Every non-empty cell is backed by a verbatim quote from the source (hover in HTML report).", lang))
    if bad:
        t.notes.append(tr("{n} cell(s) failed verification - marked ⚠.", lang).format(n=bad))
    return t


# ---------------------------------------------------------------------------
# other built-in tables
# ---------------------------------------------------------------------------

def sources_table(st: Store, plan: Plan, statuses=INCLUDED, lang: str = "en") -> Table:
    combined = plan.combined()
    t = Table(tr("Included sources", lang), [tr(c, lang) for c in ("Study", "Year", "Design", "Venue", "Weight",
                                                                    "Text", "Flags")], kind="sources")
    rows = []
    for k, s, w in st.works(status=statuses):
        rows.append((weight(w, combined), w.year or 0, [
            {"value": f"[@{k}]", "key": k, "cite": True}, str(w.year or ""), DESIGNS.get(w.design, w.design),
            w.venue[:60], f"{weight(w, combined):.2f}", st.source_text(k)[0], ", ".join(flags(w))]))
    rows.sort(key=lambda r: (-r[0], -r[1]))
    t.rows = [r[2] for r in rows]
    t.notes.append(tr("Weight = evidence weight of the design in the active domain pack(s) (0–1).", lang))
    return t


def sof_table(st: Store, plan: Plan, lang: str = "en") -> Table:
    combined = plan.combined()
    works = {k: w for k, _, w in st.works()}
    t = Table(tr("Summary of findings", lang), [tr(c, lang) for c in (
        "Claim", "Studies", "Designs", "For / against / mixed", "Consensus", "Certainty", "Key effects", "Why")],
        kind="sof")
    for c in st.claims():
        s = summarize_claim(c, st.evidence(c["id"]), works, combined)
        designs = ", ".join(f"{d}×{s.designs.count(d)}" for d in dict.fromkeys(s.designs))
        t.rows.append([f"{c['id']}: {c['text']}", str(s.n), designs,
                       f"{len(s.keys_for)} / {len(s.keys_against)} / {len(s.keys_mixed)}",
                       f"{s.label}" + (f" ({s.consensus:.0%})" if s.consensus is not None else ""),
                       s.certainty.upper(), "; ".join(s.effects[:3]), "; ".join(s.reasons)])
    t.notes.append(tr("Consensus = design-weighted share of supporting evidence. Certainty = GRADE-style: start from "
                      "the best design, capped by the reviews' own GRADE / risk-of-bias rating, downgrade for "
                      "inconsistency, imprecision, preprint-only, industry funding.", lang))
    return t


def matrix_table(st: Store, lang: str = "en") -> Table:
    claims = st.claims()
    ev = st.evidence()
    keys = list(dict.fromkeys(e["key"] for e in ev))
    t = Table(tr("Evidence matrix (claim × source)", lang), [tr("Source", lang)] + [c["id"] for c in claims],
              kind="matrix")
    sym = {"for": "✔", "against": "✘", "mixed": "≈", "neutral": "·"}
    grid = {(e["claim_id"], e["key"]): e for e in ev}
    for k in keys:
        row = [{"value": f"[@{k}]", "key": k, "cite": True}]
        for c in claims:
            e = grid.get((c["id"], k))
            if e:
                row.append({"value": sym.get(e["stance"], "?") + ("" if (e["verified"] or 0) >= 0.85 else " ⚠"),
                            "key": k, "quote": e["quote"], "loc": e["loc"], "ok": (e["verified"] or 0) >= 0.85})
            else:
                row.append("")
        t.rows.append(row)
    t.notes.append(tr("Claims: ", lang) + " | ".join(f"{c['id']} = {c['text']}" for c in claims))
    return t


def custom_table(st: Store, spec: dict) -> tuple[Table, list[str]]:
    """spec: {title, columns, rows: [{key?, cells: [str | {value, quote?, loc?, key?}]}], notes?}"""
    t = Table(spec.get("title", "Table"), spec["columns"], notes=list(spec.get("notes", [])), kind="custom")
    problems = []
    for i, r in enumerate(spec.get("rows", []), 1):
        rkey = st.resolve(r["key"]) if r.get("key") else None
        if r.get("key") and not rkey:
            problems.append(f"row {i}: unknown key {r['key']}")
        row = []
        for j, c in enumerate(r.get("cells", [])):
            if isinstance(c, str) and c.strip().startswith("@"):
                k = st.resolve(c.strip()[1:]) or c.strip()[1:]
                row.append({"value": f"[@{k}]", "key": k, "cite": True})
                continue
            if isinstance(c, dict):
                k = st.resolve(c["key"]) if c.get("key") else rkey
                if k and (c.get("quote") or any(is_significant(n) for n in extract_numbers(str(c.get("value", ""))))):
                    sc, pr = check_cell(st, k, str(c.get("value", "")), c.get("quote", ""), c.get("loc", ""))
                    problems += [f"row {i} col '{t.columns[j] if j < len(t.columns) else j}': {p}" for p in pr]
                    row.append({**c, "key": k, "ok": sc >= 0.85, "score": sc})
                else:
                    row.append(c)
            else:
                val = "" if c is None else str(c)
                nums = [n for n in extract_numbers(val) if is_significant(n)]
                if nums and rkey:
                    sc, pr = check_cell(st, rkey, val, "")
                    pr = [p for p in pr if "without a supporting quote" not in p]
                    problems += [f"row {i} col '{t.columns[j] if j < len(t.columns) else j}': {p}" for p in pr]
                    row.append({"value": val, "key": rkey, "ok": not pr})
                elif nums:
                    problems.append(f"row {i}: number {nums[0].raw} with no source (add \"key\" to the row)")
                    row.append({"value": val, "ok": False})
                else:
                    row.append(val)
        t.rows.append(row)
    return t, problems


# ---------------------------------------------------------------------------
# forest plot from extraction 'effect' cells
# ---------------------------------------------------------------------------

EFFECT_RE = re.compile(
    r"\b(OR|RR|HR|IRR|aOR|aHR|aRR|MD|SMD|RD|d|g|r)\b\s*[=:]?\s*(-?\d+(?:[.,]\d+)?)\s*[,;(\[]?\s*"
    r"(?:95\s*%\s*(?:CI|confidence interval|PU)\s*[:=,]?\s*)?(-?\d+(?:[.,]\d+)?)\s*(?:-|–|to|,|;|\.\.)\s*(-?\d+(?:[.,]\d+)?)",
    re.I)


def parse_effect(s: str):
    m = EFFECT_RE.search((s or "").replace("·", ".").replace("−", "-"))
    if not m:
        return None
    f = lambda x: float(x.replace(",", "."))  # noqa: E731
    est, lo, hi = f(m.group(2)), f(m.group(3)), f(m.group(4))
    if not (lo <= est <= hi):
        return None
    return m.group(1).upper(), est, lo, hi


def forest_svg(st: Store, field_id: str = "effect", title: str = "", lang: str = "en") -> tuple[str, int]:
    items = []
    for c in st.cells():
        if c["field"] != field_id:
            continue
        pe = parse_effect(c["value"])
        if pe:
            items.append((c["key"], *pe, (c["verified"] or 0) >= 0.85))
    if not items:
        return "", 0
    ratio = all(m in ("OR", "RR", "HR", "IRR", "AOR", "AHR", "ARR") for _, m, *_ in items)
    tx = (lambda v: math.log(v)) if ratio and all(lo > 0 for _, _, _, lo, _, _ in items) else (lambda v: v)
    null = 1.0 if ratio else 0.0
    lo_all = min(tx(lo) for _, _, _, lo, _, _ in items + [("", "", null, null, null, True)])
    hi_all = max(tx(hi) for _, _, _, _, hi, _ in items + [("", "", null, null, null, True)])
    pad = (hi_all - lo_all) * 0.08 or 0.5
    lo_all, hi_all = lo_all - pad, hi_all + pad
    W, left, right, rowh = 760, 230, 170, 26
    Hh = 50 + rowh * len(items) + 40
    X = lambda v: left + (tx(v) - lo_all) / (hi_all - lo_all) * (W - left - right)  # noqa: E731
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {Hh}" class="forest" role="img" '
           f'aria-label="{tr("Forest plot", lang)}">',
           f'<text x="{left}" y="20" class="fp-title">{H.escape(title or tr("Effect estimates (95% CI)", lang))}</text>',
           f'<line x1="{X(null):.1f}" y1="34" x2="{X(null):.1f}" y2="{Hh - 34}" class="fp-null"/>']
    for i, (k, metric, est, lo, hi, ok) in enumerate(items):
        y = 50 + i * rowh
        out.append(f'<text x="8" y="{y + 4}" class="fp-lab">@{H.escape(k[:28])}{"" if ok else " ⚠"}</text>')
        out.append(f'<line x1="{X(lo):.1f}" y1="{y}" x2="{X(hi):.1f}" y2="{y}" class="fp-ci"/>')
        out.append(f'<rect x="{X(est) - 5:.1f}" y="{y - 5}" width="10" height="10" class="fp-pt"/>')
        out.append(f'<text x="{W - right + 12}" y="{y + 4}" class="fp-val">{metric} {est:g} [{lo:g}, {hi:g}]</text>')
    ticks = [null] + ([0.25, 0.5, 2, 4] if ratio else [])
    for tv in ticks:
        if lo_all <= tx(tv) <= hi_all:
            out.append(f'<text x="{X(tv):.1f}" y="{Hh - 16}" class="fp-tick" text-anchor="middle">{tv:g}</text>')
    if ratio:
        out.append(f'<text x="{X(null) - 8:.1f}" y="{Hh - 2}" class="fp-tick" text-anchor="end">{tr("favours intervention", lang)}</text>')
        out.append(f'<text x="{X(null) + 8:.1f}" y="{Hh - 2}" class="fp-tick">{tr("favours control", lang)}</text>')
    out.append("</svg>")
    return "\n".join(out), len(items)


# ---------------------------------------------------------------------------
# renderers
# ---------------------------------------------------------------------------

def _md_escape(s: str) -> str:
    return s.replace("|", "\\|").replace("\n", " ")


def render(t: Table, fmt: str = "md", cite_fmt=None) -> str | bytes:
    cite_fmt = cite_fmt or (lambda k: f"[@{k}]")

    def txt(c):
        if isinstance(c, dict):
            if c.get("cite"):
                return cite_fmt(c["key"])
            v = str(c.get("value", ""))
            return v + ("" if c.get("ok", True) else " ⚠")
        return "" if c is None else str(c)

    if fmt == "md":
        lines = [f"**{t.title}**", "", "| " + " | ".join(_md_escape(c) for c in t.columns) + " |",
                 "|" + "|".join("---" for _ in t.columns) + "|"]
        for r in t.rows:
            lines.append("| " + " | ".join(_md_escape(txt(c)) for c in r) + " |")
        if t.notes:
            lines += [""] + [f"*{n}*" for n in t.notes]
        return "\n".join(lines) + "\n"
    if fmt == "csv":
        buf = io.StringIO()
        wr = csv.writer(buf)
        wr.writerow(t.columns + ["_quotes"])
        for r in t.rows:
            quotes = " || ".join(f"{t.columns[i]}: \"{c.get('quote')}\" ({c.get('loc', '')})"
                                 for i, c in enumerate(r) if isinstance(c, dict) and c.get("quote"))
            wr.writerow([txt(c) for c in r] + [quotes])
        return buf.getvalue()
    if fmt == "json":
        return json.dumps({"title": t.title, "columns": t.columns, "rows": t.rows, "notes": t.notes},
                          ensure_ascii=False, indent=1)
    if fmt == "latex":
        def tex(s):
            return re.sub(r"([&%$#_{}])", r"\\\1", s).replace("~", r"\textasciitilde{}").replace("⚠", "(!)")
        cols = "l" + "p{3cm}" * (len(t.columns) - 1)
        lines = [r"\begin{table}[htbp]", r"\centering\small", rf"\caption{{{tex(t.title)}}}",
                 rf"\begin{{tabular}}{{{cols}}}", r"\toprule", " & ".join(tex(c) for c in t.columns) + r" \\", r"\midrule"]
        for r in t.rows:
            lines.append(" & ".join((rf"\cite{{{c['key']}}}" if isinstance(c, dict) and c.get("cite") else tex(txt(c)))
                                    for c in r) + r" \\")
        lines += [r"\bottomrule", r"\end{tabular}"]
        if t.notes:
            lines.append(r"\par\footnotesize " + tex(" ".join(t.notes)))
        lines.append(r"\end{table}")
        return "\n".join(lines) + "\n"
    if fmt == "html":
        return table_html(t, cite_fmt)
    if fmt == "xlsx":
        try:
            import openpyxl
            from openpyxl.comments import Comment
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            raise SystemExit("xlsx needs openpyxl: pip install openpyxl (or use --format csv)")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = re.sub(r"[^\w ]", "", t.title)[:30] or "table"
        ws.append(t.columns)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in t.rows:
            ws.append([txt(c) for c in r])
            for j, c in enumerate(r):
                if isinstance(c, dict) and c.get("quote"):
                    cell = ws.cell(row=ws.max_row, column=j + 1)
                    cell.comment = Comment(f"\"{c['quote']}\" ({c.get('loc', '')})", "factual-research")
                    if not c.get("ok", True):
                        cell.fill = PatternFill("solid", fgColor="FFE0B2")
        for n in t.notes:
            ws.append([n])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    raise SystemExit(f"unknown format {fmt}")


def table_html(t: Table, cite_fmt=None) -> str:
    cite_fmt = cite_fmt or (lambda k: f"@{k}")
    out = [f'<figure class="fr-table"><figcaption>{H.escape(t.title)}</figcaption><div class="scroll"><table>',
           "<thead><tr>" + "".join(f"<th>{H.escape(c)}</th>" for c in t.columns) + "</tr></thead><tbody>"]
    for r in t.rows:
        tds = []
        for c in r:
            if isinstance(c, dict):
                if c.get("cite"):
                    tds.append(f'<td class="study"><a href="#src-{H.escape(c["key"])}">{H.escape(cite_fmt(c["key"]))}</a></td>')
                    continue
                cls = "ok" if c.get("ok", True) else "bad"
                tip = f'"{c.get("quote", "")}" - {c.get("loc", "")}' if c.get("quote") else ""
                tds.append(f'<td class="{cls}" title="{H.escape(tip)}">{H.escape(str(c.get("value", "")))}'
                           + ("" if c.get("ok", True) else ' <span class="warn">⚠</span>')
                           + ('<span class="q">❝</span>' if c.get("quote") else "") + "</td>")
            else:
                tds.append(f"<td>{H.escape('' if c is None else str(c))}</td>")
        out.append("<tr>" + "".join(tds) + "</tr>")
    out.append("</tbody></table></div>")
    for n in t.notes:
        out.append(f'<p class="note">{H.escape(n)}</p>')
    out.append("</figure>")
    return "\n".join(out)


def write(t: Table, fmt: str, out: Path | None, cite_fmt=None) -> str:
    data = render(t, fmt, cite_fmt)
    if out:
        out.parent.mkdir(parents=True, exist_ok=True)
        if isinstance(data, bytes):
            out.write_bytes(data)
        else:
            out.write_text(data, encoding="utf-8")
        return f"wrote {out}"
    return data if isinstance(data, str) else "(binary output - use --out)"


def placeholder_table(st: Store, plan: Plan, spec: list[str], lang: str | None = None
                      ) -> tuple["Table | None", list[str], str]:
    """Build the table a `<!-- fr:table ... -->` placeholder asks for, in the
    language of the final text. Returns (table or None, problems, svg-for-forest)."""
    name, args = spec[0], spec[1:]
    lang = lang or plan.lang
    if name == "extraction":
        return extraction_table(st, plan, fields=args[0].split(",") if args else None, lang=lang), [], ""
    if name == "sources":
        return sources_table(st, plan, lang=lang), [], ""
    if name == "sof":
        return sof_table(st, plan, lang=lang), [], ""
    if name == "matrix":
        return matrix_table(st, lang), [], ""
    if name == "custom" and args:
        t, pr = custom_table(st, json.loads((st.root / args[0]).read_text(encoding="utf-8")))
        return t, [f"table {args[0]}: {x}" for x in pr], ""
    if name == "forest":
        svg, n = forest_svg(st, lang=lang)
        return None, ([] if n else ["no parsable effect estimates for a forest plot"]), svg
    return None, [f"unknown table placeholder: {' '.join(spec)}"], ""
