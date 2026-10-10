"""Citation styles, bibliography exports (BibTeX / RIS / CSL-JSON) and
`render`: turn a verified draft with [@key] markers into a finished text
with in-text citations, generated tables and a reference list (md or docx)."""
from __future__ import annotations

import json
import re
from pathlib import Path

from .citations import CITE_BLOCK, TABLE_TAG, cite_items
from .model import Work, family_name
from .model import is_org as _is_org
from .project import Plan
from .store import Store

NUMERIC = {"vancouver", "ieee", "acs", "numeric"}
STYLES = ["apa", "harvard", "chicago", "vancouver", "ieee", "acs"]


def _given_initials(author: str, dots: bool = True, space: bool = True) -> str:
    a = author.strip()
    if "," in a:
        given = a.split(",", 1)[1].strip()
    else:
        parts = a.split()
        if len(parts) > 1 and re.fullmatch(r"[A-Z]{1,3}", parts[-1]):
            return ("".join(c + ("." if dots else "") + (" " if space and dots else "") for c in parts[-1])).strip()
        given = " ".join(parts[:-1])
    inits = [p[0] for p in re.split(r"[\s-]+", given) if p and p[0].isalpha()]
    sep = " " if space else ""
    return sep.join(i + ("." if dots else "") for i in inits)


def author_list(authors: list[str], style: str) -> str:
    if not authors:
        return ""
    if style in ("vancouver", "acs"):
        fmt = [a if _is_org(a) else f"{family_name(a)} {_given_initials(a, dots=False, space=False)}" for a in authors]
        return ", ".join(fmt[:6]) + (", et al." if len(fmt) > 6 else "")
    if style == "ieee":
        fmt = [a if _is_org(a) else f"{_given_initials(a)} {family_name(a)}" for a in authors]
        if len(fmt) > 6:
            return fmt[0] + " et al."
        return ", ".join(fmt[:-1]) + (", and " if len(fmt) > 2 else " and ") + fmt[-1] if len(fmt) > 1 else fmt[0]
    fmt = [a if _is_org(a) else f"{family_name(a)}, {_given_initials(a)}" for a in authors]
    if style == "apa":
        if len(fmt) > 20:
            fmt = fmt[:19] + ["... " + fmt[-1]]
        return ", ".join(fmt[:-1]) + (", & " if len(fmt) > 1 else "") + fmt[-1] if len(fmt) > 1 else fmt[0]
    if len(fmt) > 3 and style in ("harvard", "chicago"):
        return fmt[0] + " et al."
    return ", ".join(fmt[:-1]) + " and " + fmt[-1] if len(fmt) > 1 else fmt[0]


def _ids(w: Work, style: str) -> str:
    if w.doi:
        return f"https://doi.org/{w.doi}" if style in ("apa", "harvard", "chicago") else f"doi:{w.doi}"
    if w.pmid:
        return f"PMID: {w.pmid}"
    if w.arxiv:
        return f"arXiv:{w.arxiv}"
    return w.url or ""


def _pages(p: str) -> str:
    p = (p or "").strip()
    m = re.fullmatch(r"(\S+?)\s*[-–]\s*(\S+)", p)
    if m and m.group(1) == m.group(2):
        return m.group(1)
    return p


def reference(w: Work, style: str) -> str:
    au_raw = author_list(w.authors, style) or (w.venue or "Anon.")
    au = au_raw.rstrip(".")
    yr = str(w.year or "n.d.")
    title = w.title.rstrip(".")
    venue = w.venue
    vol = w.volume + (f"({w.issue})" if w.issue else "")
    pages = _pages(w.pages)
    ident = _ids(w, style)
    if w.design == "legislation":
        return f"{title}. {w.extra.get('display_address') or venue} {w.pages}. {w.url}".strip()
    if style == "apa":
        src = f"*{venue}*" + (f", *{w.volume}*" if w.volume else "") + (f"({w.issue})" if w.issue else "") + (f", {pages}" if pages else "")
        pre = " [Preprint]" if w.is_preprint else ""
        return f"{au_raw} ({yr}). {title}{pre}. {src}. {ident}".replace(". .", ".").strip()
    if style == "harvard":
        return f"{au_raw} ({yr}) '{title}', *{venue}*" + (f", {vol}" if vol else "") + (f", pp. {pages}" if pages else "") + (f". {ident}" if ident else "") + "."
    if style == "chicago":
        return f"{au}. {yr}. \"{title}.\" *{venue}*" + (f" {vol}" if vol else "") + (f": {pages}" if pages else "") + (f". {ident}" if ident else "") + "."
    if style == "vancouver":
        return f"{au}. {title}. {venue}. {yr}" + (f";{vol}" if vol else "") + (f":{pages}" if pages else "") + (f". {ident}" if ident else "") + "."
    if style == "acs":
        return f"{au}. {title}. *{venue}* **{yr}**" + (f", *{w.volume}*" if w.volume else "") + (f", {pages}" if pages else "") + (f". {ident}" if ident else "") + "."
    if style == "ieee":
        return f"{au_raw}, \"{title},\" *{venue}*" + (f", vol. {w.volume}" if w.volume else "") + (f", no. {w.issue}" if w.issue else "") + (f", pp. {pages}" if pages else "") + f", {yr}" + (f", {ident}" if ident else "") + "."
    raise SystemExit(f"unknown style {style}; choose from {', '.join(STYLES)}")


def intext_author(w: Work) -> str:
    if not w.authors:
        return (w.venue or "Anon.")[:40]
    f = family_name(w.authors[0]) if not _is_org(w.authors[0]) else w.authors[0]
    if len(w.authors) == 1:
        return f
    if len(w.authors) == 2:
        s = family_name(w.authors[1]) if not _is_org(w.authors[1]) else w.authors[1]
        return f"{f} & {s}"
    return f"{f} et al."


def _loc_fmt(loc: str, style: str) -> str:
    l = loc.strip()
    m = re.match(r"(?:pp?\.?|page|s\.|str\.)\s*(\d+(?:\s*[-–]\s*\d+)?)", l, re.I)
    if m:
        return ("pp. " if "-" in m.group(1) or "–" in m.group(1) else "p. ") + m.group(1)
    return l


# ---------------------------------------------------------------------------
# render a draft
# ---------------------------------------------------------------------------

def render_draft(st: Store, plan: Plan, text: str, style: str, lang: str | None = None,
                 table_fmt: str = "md") -> tuple[str, list[str], list[str]]:
    """Returns (rendered markdown, ordered keys, problems)."""
    from . import tables as T
    order: list[str] = []
    problems: list[str] = []
    lang = lang or plan.lang
    works: dict[str, Work] = {}
    numeric = style in NUMERIC

    def num(k):
        if k not in order:
            order.append(k)
        return order.index(k) + 1

    def cite_fmt(k):
        w = works.get(k) or st.get(k)
        if not w:
            return f"[@{k}]"
        works[k] = w
        if numeric:
            return f"[{num(k)}]"
        num(k)
        return f"({intext_author(w)}, {w.year or 'n.d.'})"

    def repl(m):
        items = []
        for key, loc, _ in cite_items(m.group(1)):
            k = st.resolve(key)
            if not k:
                problems.append(f"unknown key @{key}")
                items.append((key, loc, None))
                continue
            works[k] = works.get(k) or st.get(k)
            items.append((k, loc, works[k]))
        if numeric:
            parts = []
            for k, loc, w in items:
                n = num(k) if w else "?"
                parts.append(f"{n}" + (f", {_loc_fmt(loc, style)}" if loc else ""))
            sep = "; " if any(loc for _, loc, _ in items) else ", "
            return "[" + sep.join(parts) + "]"
        parts = []
        for k, loc, w in items:
            if not w:
                parts.append(f"@{k}?")
                continue
            num(k)
            parts.append(f"{intext_author(w)}, {w.year or 'n.d.'}" + (f", {_loc_fmt(loc, style)}" if loc else ""))
        return "(" + "; ".join(parts) + ")"

    def table_repl(m):
        try:
            t, pr, svg = T.placeholder_table(st, plan, m.group(1).strip().split(), lang)
        except Exception as e:  # noqa: BLE001
            problems.append(f"{m.group(0)}: {e}")
            return m.group(0)
        problems.extend(pr)
        if svg:
            out = st.root / "out" / "forest.svg"
            out.parent.mkdir(exist_ok=True)
            out.write_text(svg, encoding="utf-8")
            return f"![{T.tr('Forest plot', lang)}](out/forest.svg)"
        return T.render(t, "md", cite_fmt) if t else m.group(0)

    # one pass in document order, so numeric styles number by first appearance
    combo = re.compile(TABLE_TAG + "|" + CITE_BLOCK.pattern)

    def dispatch(m):
        if m.group(0).startswith("<!--"):
            return table_repl(m)
        return repl(re.match(CITE_BLOCK.pattern, m.group(0)))

    body = combo.sub(dispatch, text)
    head = "Bibliografia" if lang == "pl" else "References"
    refs = []
    keys = order if numeric else sorted(order, key=lambda k: (intext_author(works[k]).lower(), works[k].year or 0))
    for i, k in enumerate(keys, 1):
        r = reference(works[k], style)
        refs.append(f"{i}. {r}" if numeric else f"- {r}")
    rendered = body.rstrip() + f"\n\n## {head}\n\n" + "\n".join(refs) + "\n"
    return rendered, order, problems


# ---------------------------------------------------------------------------
# bibliography exports
# ---------------------------------------------------------------------------

def bibtex(entries: list[tuple[str, Work]]) -> str:
    out = []
    for k, w in entries:
        typ = {"conference_paper": "inproceedings", "book": "book", "thesis": "phdthesis", "technical_report": "techreport",
               "preprint": "misc", "dataset": "misc", "legislation": "misc", "judgment": "misc"}.get(w.design, "article")
        f = {"title": "{" + w.title + "}", "author": " and ".join(w.authors), "year": str(w.year or ""),
             ("booktitle" if typ == "inproceedings" else "journal"): w.venue, "volume": w.volume, "number": w.issue,
             "pages": w.pages.replace("-", "--"), "doi": w.doi, "publisher": w.publisher,
             "eprint": w.arxiv, "archiveprefix": "arXiv" if w.arxiv else "", "pmid": w.pmid, "url": w.url if not w.doi else ""}
        body = ",\n".join(f"  {k2} = {{{v}}}" for k2, v in f.items() if v)
        out.append(f"@{typ}{{{k},\n{body}\n}}")
    return "\n\n".join(out) + "\n"


def ris(entries: list[tuple[str, Work]]) -> str:
    out = []
    for k, w in entries:
        ty = {"conference_paper": "CPAPER", "book": "BOOK", "thesis": "THES", "technical_report": "RPRT",
              "dataset": "DATA", "legislation": "STAT", "judgment": "CASE"}.get(w.design, "JOUR")
        lines = [f"TY  - {ty}", f"ID  - {k}", f"TI  - {w.title}"] + [f"AU  - {a}" for a in w.authors]
        for tag, v in (("PY", w.year), ("JO", w.venue), ("VL", w.volume), ("IS", w.issue), ("SP", w.pages),
                       ("DO", w.doi), ("UR", w.url), ("AB", w.abstract[:2000])):
            if v:
                lines.append(f"{tag}  - {v}")
        lines.append("ER  - ")
        out.append("\n".join(lines))
    return "\n\n".join(out) + "\n"


def csl_json(entries: list[tuple[str, Work]]) -> str:
    items = []
    for k, w in entries:
        typ = {"conference_paper": "paper-conference", "book": "book", "thesis": "thesis", "technical_report": "report",
               "dataset": "dataset", "legislation": "legislation", "judgment": "legal_case", "preprint": "article"}.get(
            w.design, "article-journal")
        authors = []
        for a in w.authors:
            if "," in a:
                fam, giv = a.split(",", 1)
                authors.append({"family": fam.strip(), "given": giv.strip()})
            else:
                authors.append({"literal": a})
        it = {"id": k, "type": typ, "title": w.title, "author": authors, "container-title": w.venue,
              "volume": w.volume, "issue": w.issue, "page": w.pages, "DOI": w.doi, "URL": w.url, "PMID": w.pmid,
              "abstract": w.abstract[:2000]}
        if w.year:
            it["issued"] = {"date-parts": [[w.year]]}
        items.append({k2: v for k2, v in it.items() if v})
    return json.dumps(items, indent=1, ensure_ascii=False)


# ---------------------------------------------------------------------------
# markdown → docx (python-docx)
# ---------------------------------------------------------------------------

def md_to_docx(md: str, out: Path, base: Path):
    try:
        import docx
        from docx.shared import Inches, Pt
    except ImportError:
        raise SystemExit("docx output needs python-docx: pip install python-docx")
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.size = Pt(11)

    def add_runs(par, text):
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
            if part.startswith("**") and part.endswith("**"):
                par.add_run(part[2:-2]).bold = True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                par.add_run(part[1:-1]).italic = True
            else:
                par.add_run(part)

    lines = md.split("\n")
    i = 0
    while i < len(lines):
        s = lines[i].rstrip()
        if not s.strip():
            i += 1
            continue
        if s.startswith("#"):
            lvl = len(s) - len(s.lstrip("#"))
            d.add_heading(s.lstrip("#").strip(), level=min(lvl, 4))
        elif s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                t = d.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                t.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, val in enumerate(row):
                        cell = t.cell(r, c)
                        cell.text = ""
                        add_runs(cell.paragraphs[0], val.replace("\\|", "|"))
                        if r == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
            continue
        elif re.match(r"!\[[^\]]*\]\(([^)]+)\)", s.strip()):
            p = re.match(r"!\[[^\]]*\]\(([^)]+)\)", s.strip()).group(1)
            img = base / p
            if img.suffix.lower() == ".svg":
                d.add_paragraph(f"[figure: {p}]")
            elif img.exists():
                d.add_picture(str(img), width=Inches(6))
        elif re.match(r"^([-*+])\s+", s.strip()):
            add_runs(d.add_paragraph(style="List Bullet"), re.sub(r"^([-*+])\s+", "", s.strip()))
        elif re.match(r"^\d+[.)]\s+", s.strip()):
            add_runs(d.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", s.strip()))
        else:
            buf = [s.strip()]
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(#|\||[-*+]\s|\d+[.)]\s|!\[)", lines[i + 1].strip()):
                i += 1
                buf.append(lines[i].strip())
            add_runs(d.add_paragraph(), " ".join(buf))
        i += 1
    d.save(out)


# ---------------------------------------------------------------------------
# Obsidian / plain-markdown vault
# ---------------------------------------------------------------------------

def obsidian(st: Store, plan: Plan, out: Path, keys: list[str]) -> int:
    from .evidence import flags, summarize_claim, weight
    comb = plan.combined()
    out.mkdir(parents=True, exist_ok=True)
    (out / "sources").mkdir(exist_ok=True)
    (out / "claims").mkdir(exist_ok=True)
    ev_by_key: dict[str, list] = {}
    for e in st.evidence():
        ev_by_key.setdefault(e["key"], []).append(e)
    n = 0
    for k in keys:
        w = st.get(k)
        fm = {"key": k, "title": w.title, "year": w.year, "venue": w.venue, "doi": w.doi or None, "pmid": w.pmid or None,
              "design": w.design, "weight": weight(w, comb), "status": st.status(k), "flags": flags(w),
              "text": st.source_text(k)[0]}
        lines = ["---"] + [f"{a}: {json.dumps(b, ensure_ascii=False)}" for a, b in fm.items() if b not in (None, [], "")] + ["---", "",
                 f"# {w.title}", "", f"{reference(w, 'apa')}", ""]
        if w.abstract:
            lines += ["## Abstract", "", w.abstract, ""]
        if ev_by_key.get(k):
            lines += ["## Evidence", ""]
            for e in ev_by_key[k]:
                lines.append(f"- [[{e['claim_id']}]] **{e['stance']}** - > \"{e['quote']}\" ({e['loc'] or '?'})"
                             + (f" - {e['effect']}" if e["effect"] else ""))
            lines.append("")
        cells = st.cells(k)
        if cells:
            lines += ["## Extraction", "", "| field | value | quote |", "|---|---|---|"]
            lines += [f"| {c['field']} | {c['value']} | \"{c['quote']}\" {c['loc'] or ''} |" for c in cells]
            lines.append("")
        if w.funding or w.coi:
            lines += ["## Funding / COI", "", w.funding, w.coi, ""]
        (out / "sources" / f"{k}.md").write_text("\n".join(lines), encoding="utf-8")
        n += 1
    works = {k: w for k, _, w in st.works()}
    for c in st.claims():
        s = summarize_claim(c, st.evidence(c["id"]), works, comb)
        body = [f"# {c['id']}: {c['text']}", "", f"- consensus: {s.label}" + (f" ({s.consensus:.0%})" if s.consensus is not None else ""),
                f"- certainty: {s.certainty}", f"- why: {'; '.join(s.reasons)}", "", "## Evidence", ""]
        for e in st.evidence(c["id"]):
            body.append(f"- [[{e['key']}]] **{e['stance']}**: \"{e['quote']}\"")
        (out / "claims" / f"{c['id']}.md").write_text("\n".join(body) + "\n", encoding="utf-8")
    index = [f"# {plan.question}", "", "## Claims", ""] + [f"- [[{c['id']}]] {c['text']}" for c in st.claims()] + \
            ["", "## Sources", ""] + [f"- [[{k}]]" for k in keys]
    (out / "index.md").write_text("\n".join(index) + "\n", encoding="utf-8")
    return n
